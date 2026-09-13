"""
Utilitaire d'export de documents, partagé par TOUS les modules de Nexus.

Plutôt que chaque module réimplémente sa propre génération de PDF/Word/
Excel/PowerPoint, il appelle ces fonctions génériques avec des données
structurées simples (titres, tableaux, paragraphes). Un seul endroit à
maintenir, un seul standard visuel pour tous les documents exportés
depuis la plateforme.

Bibliothèques utilisées (à installer via requirements.txt) :
- PDF : reportlab
- Word : python-docx
- Excel : openpyxl
- PowerPoint : python-pptx

Ces exports sont couramment ~1-2 Mo par bibliothèque ; c'est le standard
de l'industrie Python pour ce besoin (aucune de ces bibliothèques ne
nécessite de service externe ou de clé d'API).
"""
import io
from dataclasses import dataclass, field


@dataclass
class ExportSection:
    """Une section de document : un titre optionnel, du texte, et/ou un tableau."""
    heading: str | None = None
    paragraphs: list[str] = field(default_factory=list)
    table_headers: list[str] | None = None
    table_rows: list[list[str]] | None = None


@dataclass
class ExportDocument:
    """Représentation générique d'un document à exporter, indépendante du format final."""
    title: str
    subtitle: str | None = None
    sections: list[ExportSection] = field(default_factory=list)


def export_to_pdf(document: ExportDocument) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm)
    styles = getSampleStyleSheet()
    story = [Paragraph(document.title, styles["Title"])]

    if document.subtitle:
        story.append(Paragraph(document.subtitle, styles["Heading3"]))
    story.append(Spacer(1, 12))

    for section in document.sections:
        if section.heading:
            story.append(Paragraph(section.heading, styles["Heading2"]))
        for paragraph in section.paragraphs:
            story.append(Paragraph(paragraph, styles["Normal"]))
            story.append(Spacer(1, 6))
        if section.table_headers and section.table_rows is not None:
            data = [section.table_headers] + section.table_rows
            table = Table(data, hAlign="LEFT")
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a1a2e")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("FONTSIZE", (0, 0), (-1, -1), 9),
                        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
                    ]
                )
            )
            story.append(table)
        story.append(Spacer(1, 16))

    doc.build(story)
    return buffer.getvalue()


def export_to_word(document: ExportDocument) -> bytes:
    from docx import Document as WordDocument

    doc = WordDocument()
    doc.add_heading(document.title, level=0)
    if document.subtitle:
        doc.add_heading(document.subtitle, level=2)

    for section in document.sections:
        if section.heading:
            doc.add_heading(section.heading, level=1)
        for paragraph in section.paragraphs:
            doc.add_paragraph(paragraph)
        if section.table_headers and section.table_rows is not None:
            table = doc.add_table(rows=1, cols=len(section.table_headers))
            table.style = "Light Grid Accent 1"
            header_cells = table.rows[0].cells
            for i, header in enumerate(section.table_headers):
                header_cells[i].text = header
            for row in section.table_rows:
                row_cells = table.add_row().cells
                for i, value in enumerate(row):
                    row_cells[i].text = str(value)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def export_to_excel(document: ExportDocument) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)

    for i, section in enumerate(document.sections):
        sheet_name = (section.heading or f"Feuille{i + 1}")[:31]
        ws = wb.create_sheet(title=sheet_name)

        row_cursor = 1
        if section.paragraphs:
            for paragraph in section.paragraphs:
                ws.cell(row=row_cursor, column=1, value=paragraph)
                row_cursor += 1
            row_cursor += 1

        if section.table_headers and section.table_rows is not None:
            header_fill = PatternFill(start_color="1A1A2E", end_color="1A1A2E", fill_type="solid")
            for col, header in enumerate(section.table_headers, start=1):
                cell = ws.cell(row=row_cursor, column=col, value=header)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
            for row_data in section.table_rows:
                row_cursor += 1
                for col, value in enumerate(row_data, start=1):
                    ws.cell(row=row_cursor, column=col, value=value)

    if not wb.sheetnames:
        wb.create_sheet(title="Résumé")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_to_powerpoint(document: ExportDocument) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = document.title
    if document.subtitle:
        title_slide.placeholders[1].text = document.subtitle

    for section in document.sections:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = section.heading or document.title

        body = slide.placeholders[1].text_frame
        body.clear()
        first = True
        for paragraph in section.paragraphs:
            p = body.paragraphs[0] if first else body.add_paragraph()
            p.text = paragraph
            p.font.size = Pt(16)
            first = False

        if section.table_headers and section.table_rows is not None:
            rows = len(section.table_rows) + 1
            cols = len(section.table_headers)
            left, top, width, height = Inches(0.5), Inches(2.2), Inches(9), Inches(0.4 * rows)
            table_shape = slide.shapes.add_table(rows, cols, left, top, width, height).table
            for col, header in enumerate(section.table_headers):
                table_shape.cell(0, col).text = header
            for r, row_data in enumerate(section.table_rows, start=1):
                for c, value in enumerate(row_data):
                    table_shape.cell(r, c).text = str(value)

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def export_to_csv(document: ExportDocument) -> bytes:
    """
    Export CSV minimal : un tableau par section. Si plusieurs sections
    contiennent des tableaux, elles sont concaténées avec une ligne de
    titre entre chacune, pour rester lisible dans Excel/LibreOffice tout
    en gardant un seul fichier.
    """
    import csv

    text_buffer = io.StringIO()
    writer = csv.writer(text_buffer)

    writer.writerow([document.title])
    if document.subtitle:
        writer.writerow([document.subtitle])

    for section in document.sections:
        writer.writerow([])
        if section.heading:
            writer.writerow([section.heading])
        for paragraph in section.paragraphs:
            writer.writerow([paragraph])
        if section.table_headers and section.table_rows is not None:
            writer.writerow(section.table_headers)
            writer.writerows(section.table_rows)

    # BOM UTF-8 ("utf-8-sig") : garantit que les accents s'affichent
    # correctement à l'ouverture directe dans Excel sous Windows.
    return text_buffer.getvalue().encode("utf-8-sig")


EXPORT_FORMATS = {
    "pdf": (export_to_pdf, "application/pdf", "pdf"),
    "docx": (export_to_word, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    "xlsx": (export_to_excel, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "pptx": (export_to_powerpoint, "application/vnd.openxmlformats-officedocument.presentationml.presentation", "pptx"),
    "csv": (export_to_csv, "text/csv", "csv"),
}


def export_document(document: ExportDocument, file_format: str) -> tuple[bytes, str, str]:
    """
    Point d'entrée unique recommandé pour tous les modules : retourne
    (contenu_binaire, type_mime, extension). Lève ValueError si le
    format demandé n'est pas supporté, à traduire en HTTP 400 côté
    routeur de chaque module.
    """
    if file_format not in EXPORT_FORMATS:
        raise ValueError(f"Format de document non supporté : {file_format}. Formats valides : {', '.join(EXPORT_FORMATS)}.")
    generator, mime_type, extension = EXPORT_FORMATS[file_format]
    return generator(document), mime_type, extension
